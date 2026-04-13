import io
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER

def rp(v):
    return f"Rp {int(v):,}".replace(',', '.')

# Export Excel 
def to_excel(df_detail, df_rekap, grand_total, ppn, grand_total_ppn, estimasi_bawah, estimasi_atas):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Detail RAB'

    BIRU   = 'FF1F3864'
    BMUDA  = 'FFD6E4F0'
    ABU    = 'FFF2F2F2'
    KUNING = 'FFFFF2CC'
    KHAKI  = 'FFBF9000'

    def bs():
        s = Side(style='thin', color='FFB0B0B0')
        return Border(left=s, right=s, top=s, bottom=s)

    def hdr(cell, bg=BIRU, fc='FFFFFFFF'):
        cell.font      = Font(bold=True, color=fc, size=10)
        cell.fill      = PatternFill('solid', fgColor=bg)
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border    = bs()

    def dat(cell, bg=ABU, align='left'):
        cell.font      = Font(size=9)
        cell.fill      = PatternFill('solid', fgColor=bg)
        cell.alignment = Alignment(horizontal=align, vertical='center')
        cell.border    = bs()

    def est(cell, bg=KUNING, fc=KHAKI, align='center'):
        cell.font      = Font(bold=True, color=fc, size=10, italic=True)
        cell.fill      = PatternFill('solid', fgColor=bg)
        cell.alignment = Alignment(horizontal=align, vertical='center')
        cell.border    = bs()

    ws.merge_cells('A1:I1')
    ws['A1'] = 'RENCANA ANGGARAN BIAYA (RAB) DETAIL — Kab. Bandung 2026'
    hdr(ws['A1'])
    ws.row_dimensions[1].height = 30

    headers = ['No', 'Ruangan', 'Pekerjaan', 'Volume', 'Satuan',
               'Harga Material', 'Harga Upah', 'Total Satuan', 'Subtotal']
    for c, h in enumerate(headers, 1):
        hdr(ws.cell(row=3, column=c, value=h))

    row_n, prev, no = 4, None, 1
    for _, r in df_detail.iterrows():
        bg = 'FFFFFFFF' if no % 2 else ABU

        if r['ruangan'] != prev:
            ws.merge_cells(f'A{row_n}:I{row_n}')
            c = ws.cell(row=row_n, column=1, value=f"  {r['ruangan'].upper()}")
            hdr(c, bg=BMUDA, fc=BIRU)
            row_n += 1
            prev, no = r['ruangan'], 1

        vals   = [no, r['ruangan'], r['pekerjaan'], r['volume'], r['satuan'],
                  int(r['harga_mat']), int(r['harga_upah']),
                  int(r['total_satuan']), int(r['biaya_total'])]
        aligns = ['center', 'left', 'left', 'center', 'center',
                  'right', 'right', 'right', 'right']

        for col, (val, al) in enumerate(zip(vals, aligns), 1):
            c = ws.cell(row=row_n, column=col, value=val)
            dat(c, bg=bg, align=al)
            if col in [6, 7, 8, 9]:
                c.number_format = '#,##0'

        row_n += 1
        no    += 1

    # Subtotal
    ws.merge_cells(f'A{row_n}:H{row_n}')
    hdr(ws.cell(row=row_n, column=1, value='SUBTOTAL (Sebelum PPN)'), bg=BMUDA, fc=BIRU)
    c = ws.cell(row=row_n, column=9, value=int(grand_total))
    hdr(c, bg=BMUDA, fc=BIRU)
    c.number_format = '#,##0'
    row_n += 1

    # PPN 11%
    ws.merge_cells(f'A{row_n}:H{row_n}')
    hdr(ws.cell(row=row_n, column=1, value='PPN 11%'), bg=BMUDA, fc=BIRU)
    c = ws.cell(row=row_n, column=9, value=int(ppn))
    hdr(c, bg=BMUDA, fc=BIRU)
    c.number_format = '#,##0'
    row_n += 1

    # Grand Total
    ws.merge_cells(f'A{row_n}:H{row_n}')
    hdr(ws.cell(row=row_n, column=1, value='GRAND TOTAL (Termasuk PPN)'))
    c = ws.cell(row=row_n, column=9, value=int(grand_total_ppn))
    hdr(c)
    c.number_format = '#,##0'
    row_n += 1

    # Rentang Estimasi AACE Class 4
    ws.merge_cells(f'A{row_n}:H{row_n}')
    label = f'RENTANG ESTIMASI  (AACE Class 4 | −15% / +20%)      {rp(estimasi_bawah)}  –  {rp(estimasi_atas)}'
    c = ws.cell(row=row_n, column=1, value=label)
    est(c, align='left')
    ws.merge_cells(f'I{row_n}:I{row_n}')
    est(ws.cell(row=row_n, column=9, value=''))
    ws.row_dimensions[row_n].height = 20

    widths = [5, 18, 25, 10, 10, 18, 18, 18, 20]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws2 = wb.create_sheet('Rekapitulasi')
    ws2.merge_cells('A1:D1')
    ws2['A1'] = 'REKAPITULASI BIAYA PER RUANGAN'
    hdr(ws2['A1'])

    for c, h in enumerate(['No', 'Nama Ruangan', 'Luas (m2)', 'Total Biaya'], 1):
        hdr(ws2.cell(row=3, column=c, value=h))

    for i, (_, r) in enumerate(df_rekap.iterrows(), 1):
        bg = 'FFFFFFFF' if i % 2 else ABU
        vals   = [i, r['ruangan'], r['luas'], int(r['total_biaya'])]
        aligns = ['center', 'left', 'center', 'right']
        for col, (v, al) in enumerate(zip(vals, aligns), 1):
            c = ws2.cell(row=i+3, column=col, value=v)
            dat(c, bg=bg, align=al)
            if col == 4:
                c.number_format = '#,##0'

    tr = len(df_rekap) + 4

    # Subtotal
    ws2.merge_cells(f'A{tr}:C{tr}')
    hdr(ws2.cell(row=tr, column=1, value='SUBTOTAL (Sebelum PPN)'), bg=BMUDA, fc=BIRU)
    c = ws2.cell(row=tr, column=4, value=int(grand_total))
    hdr(c, bg=BMUDA, fc=BIRU)
    c.number_format = '#,##0'
    tr += 1

    # PPN 11%
    ws2.merge_cells(f'A{tr}:C{tr}')
    hdr(ws2.cell(row=tr, column=1, value='PPN 11%'), bg=BMUDA, fc=BIRU)
    c = ws2.cell(row=tr, column=4, value=int(ppn))
    hdr(c, bg=BMUDA, fc=BIRU)
    c.number_format = '#,##0'
    tr += 1

    # Grand Total
    ws2.merge_cells(f'A{tr}:C{tr}')
    hdr(ws2.cell(row=tr, column=1, value='GRAND TOTAL (Termasuk PPN)'))
    c = ws2.cell(row=tr, column=4, value=int(grand_total_ppn))
    hdr(c, bg=BIRU)
    c.number_format = '#,##0'
    tr += 1

    # Rentang Estimasi AACE Class 4
    ws2.merge_cells(f'A{tr}:C{tr}')
    est(ws2.cell(row=tr, column=1, value='RENTANG ESTIMASI  (AACE Class 4 | −15% / +20%)'), align='left')
    est(ws2.cell(row=tr, column=4, value=f'{rp(estimasi_bawah)}  –  {rp(estimasi_atas)}'), align='center')
    ws2.row_dimensions[tr].height = 20

    for i, w in enumerate([5, 30, 15, 20], 1):
        ws2.column_dimensions[get_column_letter(i)].width = w

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# Export PDF
def to_pdf(df_detail, df_rekap, grand_total, ppn, grand_total_ppn, estimasi_bawah, estimasi_atas):
    buf    = io.BytesIO()
    doc    = SimpleDocTemplate(buf, pagesize=A4,
                               leftMargin=1*cm, rightMargin=1*cm,
                               topMargin=1.5*cm, bottomMargin=1.5*cm)
    styles = getSampleStyleSheet()
    BIRU   = colors.HexColor('#1F3864')
    BMUDA  = colors.HexColor('#D6E4F0')
    ABU    = colors.HexColor('#F2F2F2')
    KUNING = colors.HexColor('#FFF2CC')
    KHAKI  = colors.HexColor('#BF9000')

    story = [
        Paragraph('RENCANA ANGGARAN BIAYA (RAB)', ParagraphStyle(
            'j', parent=styles['Title'], fontSize=14, textColor=BIRU,
            alignment=TA_CENTER, spaceAfter=2)),
        Paragraph('Kabupaten Bandung — Perbup No.55 Tahun 2025', ParagraphStyle(
            's', parent=styles['Normal'], fontSize=8, textColor=colors.grey,
            alignment=TA_CENTER, spaceAfter=10)),
    ]

    for ruangan in df_detail['ruangan'].unique():
        story.append(Paragraph(f'RUANGAN: {ruangan.upper()}', ParagraphStyle(
            'r', parent=styles['Normal'], fontSize=9,
            fontName='Helvetica-Bold', spaceBefore=10, spaceAfter=4)))

        df_r = df_detail[df_detail['ruangan'] == ruangan]

        data = [['No', 'Pekerjaan', 'Vol', 'Sat',
                 'Harga Mat', 'Harga Upah', 'Subtotal']]

        for i, (_, row) in enumerate(df_r.iterrows(), 1):
            data.append([
                str(i), row['pekerjaan'],
                f"{row['volume']:.2f}", row['satuan'],
                f"{int(row['harga_mat']):,}".replace(',', '.'),
                f"{int(row['harga_upah']):,}".replace(',', '.'),
                f"{int(row['biaya_total']):,}".replace(',', '.'),
            ])

        data.append(['', 'Total Biaya Ruang', '', '', '', '',
                     f"{int(df_r['biaya_total'].sum()):,}".replace(',', '.')])

        t = Table(data, colWidths=[0.7*cm, 5.0*cm, 1.2*cm, 1.1*cm,
                                   3.5*cm, 3.5*cm, 4.0*cm])
        t.setStyle(TableStyle([
            ('BACKGROUND',     (0, 0),  (-1, 0),  BIRU),
            ('TEXTCOLOR',      (0, 0),  (-1, 0),  colors.white),
            ('FONTNAME',       (0, 0),  (-1, 0),  'Helvetica-Bold'),
            ('FONTSIZE',       (0, 0),  (-1, -1), 7),
            ('ROWBACKGROUNDS', (0, 1),  (-1, -2), [colors.white, ABU]),
            ('BACKGROUND',     (0, -1), (-1, -1), BMUDA),
            ('FONTNAME',       (0, -1), (-1, -1), 'Helvetica-Bold'),
            ('ALIGN',          (4, 0),  (-1, -1), 'RIGHT'),
            ('ALIGN',          (0, 0),  (0, -1),  'CENTER'),
            ('GRID',           (0, 0),  (-1, -1), 0.2, colors.grey),
            ('VALIGN',         (0, 0),  (-1, -1), 'MIDDLE'),
        ]))
        story.append(t)
        story.append(Spacer(1, 0.4*cm))

    # Rekapitulasi 
    story.append(Spacer(1, 0.8*cm))
    story.append(Paragraph('REKAPITULASI TOTAL BIAYA', ParagraphStyle(
        'rek', parent=styles['Normal'], fontSize=11,
        fontName='Helvetica-Bold', textColor=BIRU, spaceAfter=6)))

    rekap_data = [['No', 'Uraian Ruangan', 'Luas (m2)', 'Total Biaya']]
    for i, (_, r) in enumerate(df_rekap.iterrows(), 1):
        rekap_data.append([
            str(i), r['ruangan'],
            f"{r['luas']:.2f}",
            f"Rp {int(r['total_biaya']):,}".replace(',', '.'),
        ])

    rekap_data.append(['', 'SUBTOTAL (Sebelum PPN)', '',
                       f"Rp {int(grand_total):,}".replace(',', '.')])
    rekap_data.append(['', 'PPN 11%', '',
                       f"Rp {int(ppn):,}".replace(',', '.')])
    rekap_data.append(['', 'GRAND TOTAL (Termasuk PPN)', '',
                       f"Rp {int(grand_total_ppn):,}".replace(',', '.')])
    rekap_data.append(['', 'RENTANG ESTIMASI  (AACE Class 4 | −15% / +20%)', '',
                       f"{rp(estimasi_bawah)}  –  {rp(estimasi_atas)}"])

    t2 = Table(rekap_data, colWidths=[1*cm, 8.5*cm, 2.5*cm, 7*cm])
    t2.setStyle(TableStyle([
        ('BACKGROUND',   (0, 0),  (-1, 0),  BIRU),
        ('TEXTCOLOR',    (0, 0),  (-1, 0),  colors.white),
        ('FONTNAME',     (0, 0),  (-1, 0),  'Helvetica-Bold'),
        ('FONTSIZE',     (0, 0),  (-1, -1), 9),
        ('BACKGROUND',   (0, -4), (-1, -4), BMUDA),
        ('FONTNAME',     (0, -4), (-1, -4), 'Helvetica-Bold'),
        ('BACKGROUND',   (0, -3), (-1, -2), BIRU),
        ('TEXTCOLOR',    (0, -3), (-1, -2), colors.white),
        ('FONTNAME',     (0, -3), (-1, -2), 'Helvetica-Bold'),
        ('BACKGROUND',   (0, -1), (-1, -1), KUNING),
        ('TEXTCOLOR',    (0, -1), (-1, -1), KHAKI),
        ('FONTNAME',     (0, -1), (-1, -1), 'Helvetica-BoldOblique'),
        ('ALIGN',        (-1, 0), (-1, -1), 'RIGHT'),
        ('GRID',         (0, 0),  (-1, -1), 0.5, colors.black),
        ('TOPPADDING',   (0, 0),  (-1, -1), 6),
        ('BOTTOMPADDING',(0, 0),  (-1, -1), 6),
    ]))
    story.append(t2)

    # Catatan kaki
    story.append(Spacer(1, 0.5*cm))
    story.append(Paragraph(
        '* Rentang estimasi berdasarkan AACE International Class 4 Estimate (−15% / +20%). '
        'Angka ini merupakan perkiraan awal untuk tahap perencanaan dan belum bersifat final.',
        ParagraphStyle('note', parent=styles['Normal'], fontSize=7,
                       textColor=colors.grey, spaceBefore=4)
    ))

    doc.build(story)
    buf.seek(0)
    return buf